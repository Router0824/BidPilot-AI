import io
import asyncio
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
from docx import Document as WordDocument
from fastapi import HTTPException
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker

from app.main import app
from app.core.auth import require_auth
from app.core.database import Base, get_db
from app.core.config import settings
from app.domain.models import (Project, OutlineSection, DraftVersion, SectionProtection,
    KnowledgeChunk, ReviewRun, ReviewFinding, Requirement, Document, ProjectMember)
from app.application.editing_service import write_version, set_protection
from app.application.review_export_service import ReviewService, FixerService, ExportService
from app.application.knowledge_service import KnowledgeIndexService
from app.agents import drafting_agent


class ProductivityTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.engine = create_async_engine(f"sqlite+aiosqlite:///{self.tmp.name}/test.db")
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.user = {"id": "admin", "role": "admin"}

        async def db_override():
            async with self.sessions() as db:
                try:
                    yield db
                    await db.commit()
                except Exception:
                    await db.rollback()
                    raise

        app.dependency_overrides[get_db] = db_override
        app.dependency_overrides[require_auth] = lambda: self.user
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")
        self.old_upload = settings.UPLOAD_DIR
        settings.UPLOAD_DIR = self.tmp.name
        async with self.sessions() as db:
            db.add_all([Project(id="p", name="实用版测试", owner_id="admin"),
                        Project(id="other", name="其他项目", owner_id="someone"),
                        OutlineSection(id="s", project_id="p", title="技术方案", sort_order=1),
                        OutlineSection(id="other-s", project_id="other", title="其他章节")])
            await db.commit()

    async def asyncTearDown(self):
        await self.client.aclose()
        app.dependency_overrides.clear()
        settings.UPLOAD_DIR = self.old_upload
        await self.engine.dispose()
        self.tmp.cleanup()

    async def save(self, content="人工正文", version=None, **extra):
        return await self.client.put("/api/v1/projects/p/outline/sections/s/editor", json={
            "content": content, "expected_version_id": version, **extra})

    async def test_manual_save_survives_reload_and_protects(self):
        response = await self.save()
        self.assertEqual(response.status_code, 200, response.text)
        current = (await self.client.get("/api/v1/projects/p/outline/sections/s/editor")).json()["data"]
        self.assertEqual(current["content"], "人工正文")
        self.assertTrue(current["protected"])
        async with self.sessions() as db:
            with patch.object(drafting_agent, "_generate_section_content", new=AsyncMock()) as generation:
                result = await drafting_agent.generate_draft("p", "s", db)
                self.assertTrue(result["skipped"])
                generation.assert_not_called()

    async def test_stale_version_rejected_without_losing_history(self):
        first = (await self.save()).json()["data"]["version_id"]
        second = await self.save("第二版", first)
        self.assertEqual(second.status_code, 200)
        conflict = await self.save("过时编辑", first)
        self.assertEqual(conflict.status_code, 409)
        async with self.sessions() as db:
            history = list((await db.scalars(select(DraftVersion))).all())
            self.assertEqual({d.content for d in history}, {"人工正文", "第二版"})

    async def test_cross_project_section_is_rejected(self):
        response = await self.client.put("/api/v1/projects/p/outline/sections/other-s/editor",
            json={"content": "bad", "expected_version_id": None})
        self.assertEqual(response.status_code, 404)

    async def test_nonmember_cannot_read_editor(self):
        self.user = {"id": "writer", "role": "writer"}
        response = await self.client.get("/api/v1/projects/p/outline/sections/s/editor")
        self.assertEqual(response.status_code, 403)

    async def test_generated_draft_has_current_version_and_outline_is_preserved(self):
        async with self.sessions() as db:
            with patch.object(drafting_agent, "_generate_section_content", new=AsyncMock(return_value="模型正文")):
                result = await drafting_agent.generate_draft("p", "s", db)
            section = await db.get(OutlineSection, "s")
            self.assertIsNotNone(result["draft_id"])
            self.assertEqual(section.current_version_id, result["draft_id"])
            await drafting_agent.generate_outline("p", db)
            count = await db.scalar(select(func.count()).select_from(OutlineSection).where(OutlineSection.project_id == "p"))
            self.assertEqual(count, 1)

    async def test_late_generation_cannot_overwrite_manual_save(self):
        async with self.sessions() as db:
            section = await db.get(OutlineSection, "s")
            draft = await write_version(db, section, "人工正文", [], None, actor="admin", protect=True)
            with self.assertRaises(HTTPException):
                await write_version(db, section, "迟到的生成", [], None)
            self.assertEqual(section.current_version_id, draft.id)

    async def test_expired_material_not_retrieved_or_saved_as_citation(self):
        async with self.sessions() as db:
            db.add(KnowledgeChunk(id="expired", material_name="过期资质", content="技术能力",
                is_audited=True, valid_until=datetime.now(timezone.utc) - timedelta(days=1)))
            db.add(KnowledgeChunk(id="valid", material_name="有效资质", content="技术能力", is_audited=True))
            await db.commit()
            found = await KnowledgeIndexService(db).retrieve("技术能力")
            self.assertEqual([row["id"] for row in found], ["valid"])
        self.assertEqual((await self.save(knowledge_ids=["expired"])).status_code, 422)
        self.assertEqual((await self.save(knowledge_ids=["missing"])).status_code, 422)

    async def test_review_findings_reload_and_require_reasons(self):
        await self.save()
        response = await self.client.post("/api/v1/projects/p/reviews")
        self.assertEqual(response.status_code, 200)
        findings = (await self.client.get("/api/v1/projects/p/review-tasks")).json()["data"]
        self.assertTrue(findings)
        self.assertEqual(findings[0]["section_id"], "s")
        url = f"/api/v1/projects/p/reviews/findings/{findings[0]['id']}"
        self.assertEqual((await self.client.patch(url, params={"status": "ignored"})).status_code, 422)
        self.assertEqual((await self.client.patch(url, params={"status": "ignored", "ignore_reason": "已人工核实"})).status_code, 200)
        latest = (await self.client.get("/api/v1/projects/p/review-tasks")).json()["data"]
        self.assertEqual(latest[0]["ignore_reason"], "已人工核实")

    async def test_high_risk_confirmation_requires_reviewer(self):
        async with self.sessions() as db:
            db.add(ProjectMember(project_id="p", user_id="writer"))
            db.add(ReviewRun(id="r", project_id="p", review_type="full"))
            await db.flush()
            db.add(ReviewFinding(id="high", review_run_id="r", risk_level="high", status="open"))
            await db.commit()
        self.user = {"id": "writer", "role": "writer"}
        response = await self.client.patch("/api/v1/projects/p/reviews/findings/high",
            params={"status": "resolved", "ignore_reason": "test"})
        self.assertEqual(response.status_code, 403)

    async def test_fixer_creates_version_preserves_original_and_honors_protection(self):
        async with self.sessions() as db:
            section = await db.get(OutlineSection, "s")
            original = await write_version(db, section, "原文", [], None)
            db.add(ReviewRun(id="r", project_id="p", review_type="full"))
            await db.flush()
            db.add(ReviewFinding(id="f", review_run_id="r", risk_level="low", finding_type="citation_missing", location="section:s"))
            await db.flush()
            first = await FixerService(db).fix_issue("p", "f", self.user)
            self.assertEqual(first.status, "applied")
            self.assertEqual(original.content, "原文")
            self.assertNotEqual(section.current_version_id, original.id)
            await set_protection(db, "s", True, "admin")
            second = await FixerService(db).fix_issue("p", "f", self.user)
            self.assertEqual(second.status, "manual_required")

    async def test_fixer_stops_after_two_applied_attempts(self):
        from app.domain.models import FixAttempt
        async with self.sessions() as db:
            section = await db.get(OutlineSection, "s")
            await write_version(db, section, "原文", [], None)
            db.add(ReviewRun(id="r", project_id="p", review_type="full"))
            await db.flush()
            db.add(ReviewFinding(id="f", review_run_id="r", risk_level="low", finding_type="citation_missing", location="section:s"))
            await db.flush()
            db.add_all([FixAttempt(project_id="p", issue_id="f", section_id="s", status="applied", attempt_no=i) for i in (1, 2)])
            await db.flush()
            attempt = await FixerService(db).fix_issue("p", "f", self.user)
            self.assertEqual(attempt.status, "manual_required")
            self.assertEqual(attempt.attempt_no, 3)

    async def test_workbench_uses_actual_data_and_blocks_final_export(self):
        async with self.sessions() as db:
            db.add(Requirement(id="req", project_id="p", requirement_text="需资质", requirement_type="qualification", risk_level="high"))
            await db.commit()
        data = (await self.client.get("/api/v1/projects/p/workbench")).json()["data"]
        self.assertEqual(data["missing_materials"], 1)
        self.assertEqual(data["completed_sections"], 0)
        self.assertFalse(data["export_ready"])
        response = await self.client.post("/api/v1/projects/p/delivery/download", json={"mode": "final"})
        self.assertEqual(response.status_code, 409)

    async def test_draft_download_contains_tables_and_fields(self):
        await self.save("| 阶段 | 交付物 |\n| --- | --- |\n| 设计 | 设计文档 |")
        response = await self.client.post("/api/v1/projects/p/delivery/download", json={"company_name": "测试企业"})
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else "")
        document = WordDocument(io.BytesIO(response.content))
        self.assertEqual(document.tables[0].cell(1, 1).text, "设计文档")
        self.assertIn("TOC", document.element.xml)
        self.assertIn("PAGE", document.sections[-1].footer._element.xml)
        self.assertIn("测试企业", "\n".join(p.text for p in document.paragraphs))

    async def test_template_upload_rejects_non_docx_and_preserves_template(self):
        bad = await self.client.post("/api/v1/projects/p/delivery/template", files={"file": ("bad.docx", b"invalid")})
        self.assertEqual(bad.status_code, 422)
        doc = WordDocument(); doc.add_paragraph("企业模板封面")
        buffer = io.BytesIO(); doc.save(buffer)
        upload = await self.client.post("/api/v1/projects/p/delivery/template", files={"file": ("企业.docx", buffer.getvalue())})
        self.assertEqual(upload.status_code, 200)
        response = await self.client.post("/api/v1/projects/p/delivery/download", json={"mode": "draft"})
        exported = WordDocument(io.BytesIO(response.content))
        self.assertEqual(exported.paragraphs[0].text, "企业模板封面")

    async def test_new_tables_initialization_preserves_existing_data(self):
        from app.domain.models import MaterialMetadata, ExportProfile
        async with self.engine.begin() as conn:
            for model in (SectionProtection, MaterialMetadata, ExportProfile):
                await conn.run_sync(lambda connection, table=model.__table__: table.drop(connection))
            await conn.run_sync(Base.metadata.create_all)
            await conn.run_sync(Base.metadata.create_all)
        async with self.sessions() as db:
            self.assertEqual((await db.get(Project, "p")).name, "实用版测试")

    async def test_concurrent_saves_have_only_one_winner(self):
        responses = await asyncio.gather(self.save("甲的版本"), self.save("乙的版本"))
        self.assertEqual(sorted(r.status_code for r in responses), [200, 409])
        async with self.sessions() as db:
            self.assertEqual(await db.scalar(select(func.count()).select_from(DraftVersion)), 1)

    async def test_complete_delivery_then_edit_requires_reconfirmation(self):
        async with self.sessions() as db:
            db.add(Document(project_id="p", name="招标文件.txt", parse_status="completed"))
            db.add(KnowledgeChunk(id="valid", material_name="企业材料", source_page=1, content="技术能力", is_audited=True))
            db.add(Requirement(id="req", project_id="p", requirement_text="技术能力", requirement_type="technical", response_section_id="s"))
            await db.commit()
        first = await self.save("企业技术能力说明", knowledge_ids=["valid"])
        self.assertEqual(first.status_code, 200)
        response = await self.client.patch("/api/v1/projects/p/requirements/req", json={"status": "responded"})
        self.assertEqual(response.status_code, 200, response.text)
        review = await self.client.post("/api/v1/projects/p/reviews")
        self.assertEqual(review.status_code, 200)
        result = (await self.client.get("/api/v1/projects/p/workbench")).json()["data"]
        self.assertTrue(result["export_ready"], result["tasks"])
        final = await self.client.post("/api/v1/projects/p/delivery/download", json={"mode": "final"})
        self.assertEqual(final.status_code, 200)
        await self.save("修改后的技术能力", first.json()["data"]["version_id"])
        after = (await self.client.get("/api/v1/projects/p/workbench")).json()["data"]
        self.assertFalse(after["export_ready"])
        self.assertIn("review-stale", [t["id"] for t in after["tasks"]])

    async def test_generation_mock_path_preserves_manual_version_after_unlock(self):
        first = (await self.save()).json()["data"]["version_id"]
        await self.client.put("/api/v1/projects/p/outline/sections/s/protection", json={"protected": False})
        from app.agents import MockLLMGateway
        with patch.object(drafting_agent, "llm", MockLLMGateway()):
            response = await self.client.post("/api/v1/projects/p/outline/sections/s/draft", json={})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotEqual(response.json()["data"]["draft_id"], first)
        async with self.sessions() as db:
            self.assertEqual((await db.get(DraftVersion, first)).content, "人工正文")

    async def test_sse_releases_database_before_yielding(self):
        from contextlib import asynccontextmanager
        from app.api.v1.workflows import stream_workflow_status
        from app.core.auth import create_access_token
        active = []

        @asynccontextmanager
        async def tracked_session():
            active.append(True)
            async with self.sessions() as db:
                try:
                    yield db
                finally:
                    active.pop()

        request = AsyncMock()
        request.is_disconnected.return_value = False
        with patch("app.core.database.async_session", tracked_session):
            response = await stream_workflow_status("p", request, create_access_token({"sub": "admin"}))
            event = await anext(response.body_iterator)
            self.assertEqual(event["event"], "workflow.status.changed")
            self.assertEqual(active, [])
            await response.body_iterator.aclose()

    async def test_template_fields_across_runs_tables_headers_and_body_position(self):
        document = WordDocument()
        paragraph = document.add_paragraph()
        paragraph.add_run('项目：{{pro').bold = True
        paragraph.add_run('ject_name}} / {{company_name}}')
        document.add_table(rows=1, cols=1).cell(0, 0).text = '{{company_name}}'
        document.sections[0].header.paragraphs[0].text = '{{project_name}}'
        document.add_paragraph('{{content}}')
        document.add_paragraph('固定附件位置')
        buffer = io.BytesIO(); document.save(buffer)
        upload = await self.client.post('/api/v1/projects/p/delivery/template', files={'file': ('template.docx', buffer.getvalue())})
        self.assertEqual(upload.status_code, 200, upload.text)
        self.assertEqual(upload.json()['data']['template_info']['insertion'], 'placeholder')
        await self.save('正文插入验收')
        response = await self.client.post('/api/v1/projects/p/delivery/download', json={'company_name': '测试单位'})
        self.assertEqual(response.status_code, 200)
        exported = WordDocument(io.BytesIO(response.content))
        self.assertEqual(exported.paragraphs[0].text, '项目：实用版测试 / 测试单位')
        self.assertTrue(exported.paragraphs[0].runs[0].bold)
        self.assertEqual(exported.tables[0].cell(0, 0).text, '测试单位')
        self.assertEqual(exported.sections[0].header.paragraphs[0].text, '实用版测试')
        text = '\n'.join(p.text for p in exported.paragraphs)
        self.assertNotIn('{{', text)
        self.assertLess(text.index('正文插入验收'), text.index('固定附件位置'))

    async def test_invalid_template_does_not_replace_existing_one(self):
        sample = await self.client.get('/api/v1/projects/p/delivery/template/sample')
        self.assertEqual(sample.status_code, 200)
        await self.client.post('/api/v1/projects/p/delivery/template', files={'file': ('valid.docx', sample.content)})
        for mode in ('duplicate', 'table', 'unknown'):
            document = WordDocument()
            if mode == 'duplicate':
                document.add_paragraph('{{content}}'); document.add_paragraph('{{content}}')
            elif mode == 'table':
                document.add_table(rows=1, cols=1).cell(0, 0).text = '{{content}}'
            else:
                document.add_paragraph('{{price}}')
            output = io.BytesIO(); document.save(output)
            response = await self.client.post('/api/v1/projects/p/delivery/template', files={'file': ('bad.docx', output.getvalue())})
            self.assertEqual(response.status_code, 422, mode)
        profile = (await self.client.get('/api/v1/projects/p/delivery/profile')).json()['data']
        self.assertEqual(profile['template_name'], 'valid.docx')

    async def test_template_requires_company_and_can_be_reset(self):
        sample = await self.client.get('/api/v1/projects/p/delivery/template/sample')
        await self.client.post('/api/v1/projects/p/delivery/template', files={'file': ('sample.docx', sample.content)})
        missing = await self.client.post('/api/v1/projects/p/delivery/download', json={})
        self.assertEqual(missing.status_code, 422)
        reset = await self.client.delete('/api/v1/projects/p/delivery/template')
        self.assertEqual(reset.status_code, 200)
        default = await self.client.post('/api/v1/projects/p/delivery/download', json={})
        self.assertEqual(default.status_code, 200)


if __name__ == "__main__":
    unittest.main()
