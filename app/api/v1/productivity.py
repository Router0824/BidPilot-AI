from typing import Literal
from pathlib import Path
from uuid import uuid4
from docx.opc.exceptions import PackageNotFoundError
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.database import get_db
from app.core.auth import require_auth
from app.core.config import settings
from app.domain.models import (Project, ProjectMember, OutlineSection, DraftVersion,
    SectionProtection, ReviewRun, ReviewFinding, ExportProfile)
from app.application.editing_service import editable_section, write_version, set_protection
from app.application.enterprise_service import EnterpriseService
from app.application.workbench_service import workbench
from app.application.review_export_service import ExportService
from app.schemas import APIResponse


async def project_access(project_id: str, db=Depends(get_db), user=Depends(require_auth)):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "项目不存在")
    member = await db.scalar(select(ProjectMember.id).where(
        ProjectMember.project_id == project_id, ProjectMember.user_id == user["id"]))
    if user["role"] not in ("admin", "project_admin") and project.owner_id != user["id"] and not member:
        raise HTTPException(403, "您不是该项目成员")


router = APIRouter(prefix="/projects/{project_id}", tags=["workbench"], dependencies=[Depends(project_access)])


class SaveDraft(BaseModel):
    content: str = Field(max_length=200000)
    expected_version_id: str | None
    knowledge_ids: list[str] | None = Field(default=None, max_length=50)


class ProtectionUpdate(BaseModel):
    protected: bool


class DeliveryRequest(BaseModel):
    company_name: str = Field(default="", max_length=255)
    mode: Literal["draft", "final"] = "draft"


@router.get("/workbench")
async def get_workbench(project_id: str, db=Depends(get_db)):
    return APIResponse(data=await workbench(db, project_id))


@router.get("/outline/sections/{section_id}/editor")
async def editor(project_id: str, section_id: str, db=Depends(get_db)):
    section = await db.get(OutlineSection, section_id)
    if not section or section.project_id != project_id:
        raise HTTPException(404, "章节不存在")
    draft = await db.get(DraftVersion, section.current_version_id) if section.current_version_id else None
    protection = await db.get(SectionProtection, section_id)
    return APIResponse(data={"version_id": section.current_version_id, "content": draft.content if draft else "",
        "citations": draft.citations if draft else [], "protected": bool(protection and protection.protected)})


@router.put("/outline/sections/{section_id}/editor")
async def save_editor(project_id: str, section_id: str, data: SaveDraft,
                      db=Depends(get_db), user=Depends(require_auth)):
    section = await editable_section(db, project_id, section_id, user)
    old = await db.get(DraftVersion, section.current_version_id) if section.current_version_id else None
    citations = old.citations if old else []
    if data.knowledge_ids is not None:
        from app.domain.models import KnowledgeChunk
        from datetime import datetime, timezone
        citations = []
        for chunk_id in dict.fromkeys(data.knowledge_ids):
            chunk = await db.get(KnowledgeChunk, chunk_id)
            if not chunk or not chunk.is_audited or chunk.is_expired or (
                chunk.valid_until and chunk.valid_until.replace(tzinfo=timezone.utc) <= datetime.now(timezone.utc)):
                raise HTTPException(422, "引用材料不存在、未审核或已过期，请重新选择")
            citations.append({"chunk_id": chunk.id, "source": chunk.material_name, "page": chunk.source_page,
                "snippet": (chunk.content or "")[:300], "status": "verified", "version": chunk.document_version})
    draft = await write_version(db, section, data.content, citations,
                                data.expected_version_id, actor=user["id"], protect=True)
    await EnterpriseService(db).audit("draft", project_id, "manual_save", user,
        {"version_id": data.expected_version_id}, {"version_id": draft.id, "section_id": section_id})
    from app.application.evidence_service import EvidenceGraphService
    await EvidenceGraphService(db).rebuild_project_links(project_id)
    return APIResponse(data={"version_id": draft.id, "protected": True})


@router.put("/outline/sections/{section_id}/protection")
async def protect_editor(project_id: str, section_id: str, data: ProtectionUpdate,
                         db=Depends(get_db), user=Depends(require_auth)):
    await editable_section(db, project_id, section_id, user)
    await set_protection(db, section_id, data.protected, user["id"])
    await EnterpriseService(db).audit("draft", project_id, "protection", user, None,
        {"section_id": section_id, "protected": data.protected})
    return APIResponse(data={"protected": data.protected})


@router.get("/review-tasks")
async def review_tasks(project_id: str, db=Depends(get_db)):
    run = await db.scalar(select(ReviewRun).where(ReviewRun.project_id == project_id).order_by(
        ReviewRun.created_at.desc()).limit(1))
    findings = (await db.scalars(select(ReviewFinding).where(ReviewFinding.review_run_id == run.id))).all() if run else []
    sections = (await db.scalars(select(OutlineSection).where(OutlineSection.project_id == project_id))).all()
    return APIResponse(data=[{
        "id": f.id, "finding_type": f.finding_type, "risk_level": f.risk_level, "description": f.description,
        "suggestion": f.suggestion, "status": f.status, "ignore_reason": f.ignore_reason,
        "section_id": next((s.id for s in sections if f.location in (s.title, f"section:{s.id}")), None),
        "auto_fix_allowed": f.risk_level == "low" and f.finding_type == "citation_missing",
    } for f in findings])


@router.post("/delivery/template")
async def upload_template(project_id: str, file: UploadFile = File(...), db=Depends(get_db), user=Depends(require_auth)):
    if user["role"] not in ("admin", "project_admin"):
        raise HTTPException(403, "仅项目管理员可修改模板")
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024 or not (file.filename or "").lower().endswith(".docx"):
        raise HTTPException(422, "请上传不超过 5 MB 的 DOCX 模板")
    import io, zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("template too large")
            if any("vbaProject" in item.filename for item in archive.infolist()):
                raise ValueError("macro template")
        from docx import Document
        document = Document(io.BytesIO(content))
    except Exception:
        raise HTTPException(422, "模板不是有效的无宏 Word 文档")
    from app.application.word_template import inspect_template
    try:
        template_info = inspect_template(document)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    path = Path(settings.UPLOAD_DIR) / project_id / "templates" / f"{uuid4().hex}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    profile = await db.get(ExportProfile, project_id)
    if profile is None:
        profile = ExportProfile(project_id=project_id)
        db.add(profile)
    profile.template_path, profile.template_name = str(path), file.filename
    await db.flush()
    await EnterpriseService(db).audit("export_template", project_id, "upload", user, None,
        {"template_name": profile.template_name, **template_info})
    return APIResponse(data={"template_name": profile.template_name, "template_info": template_info})


@router.get("/delivery/template/sample")
async def sample_template(project_id: str):
    import io
    from docx import Document
    doc = Document()
    doc.add_paragraph("{{project_name}}", style="Title")
    doc.add_paragraph("投标单位：{{company_name}}")
    doc.add_page_break()
    doc.add_paragraph("{{content}}")
    doc.add_page_break()
    doc.add_heading("附件", level=1)
    output = io.BytesIO()
    doc.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": 'attachment; filename="bidpilot-template.docx"'})


@router.delete("/delivery/template")
async def reset_template(project_id: str, db=Depends(get_db), user=Depends(require_auth)):
    if user["role"] not in ("admin", "project_admin"):
        raise HTTPException(403, "仅项目管理员可修改模板")
    profile = await db.get(ExportProfile, project_id)
    if profile:
        profile.template_path = None
        profile.template_name = None
    await EnterpriseService(db).audit("export_template", project_id, "reset", user)
    return APIResponse(data={"template_name": None})


@router.get("/delivery/profile")
async def delivery_profile(project_id: str, db=Depends(get_db)):
    profile = await db.get(ExportProfile, project_id)
    info = None
    if profile and profile.template_path:
        from docx import Document
        from app.application.word_template import inspect_template
        try:
            info = inspect_template(Document(profile.template_path))
        except (ValueError, OSError, PackageNotFoundError):
            info = {"error": "模板不可用，请重新上传或恢复默认模板"}
    return APIResponse(data={"company_name": profile.company_name if profile else "",
                             "template_info": info,
                             "template_name": profile.template_name if profile else None})


@router.post("/delivery/download")
async def download_delivery(project_id: str, data: DeliveryRequest, db=Depends(get_db), user=Depends(require_auth)):
    check = await workbench(db, project_id)
    if data.mode == "final" and not check["export_ready"]:
        raise HTTPException(409, "仍有未完成事项，请处理后导出正式稿，或选择工作草稿")
    profile = await db.get(ExportProfile, project_id)
    project = await db.get(Project, project_id)
    svc = ExportService(db)
    directory = Path(settings.UPLOAD_DIR) / project_id / "exports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{uuid4().hex}.docx"
    sections = list((await db.scalars(select(OutlineSection).where(
        OutlineSection.project_id == project_id).order_by(OutlineSection.sort_order))).all())
    try:
        await svc._write_full_docx(str(path), sections, title=project.name, company=data.company_name,
            template=profile.template_path if profile else None, draft_mode=data.mode == "draft")
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    except (OSError, PackageNotFoundError):
        raise HTTPException(422, "模板文件不可用，请重新上传或恢复默认模板")
    if profile is None:
        profile = ExportProfile(project_id=project_id)
        db.add(profile)
    profile.company_name = data.company_name
    await EnterpriseService(db).audit("export", project_id, "download", user, None,
        {"mode": data.mode, "file": path.name, "pending_tasks": len(check["tasks"])})
    return FileResponse(str(path), filename=f"{project.name[:80]}-{'工作草稿' if data.mode == 'draft' else '技术标'}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
